"""Generic re-import of a previously-exported Recon OS report (any module).

Every module's `build_report()` writes one header row + data rows per
sheet in a consistent shape, so ONE importer works for all seven modules
(and reasonably for the original tools' own exports, which follow the same
header-row-then-data-rows convention) - it doesn't re-run any module's
parse/reconcile logic, it just reconstructs a completed Run's persisted
RunResults directly from the workbook, for archival/legacy access
(PRD §15 open question: legacy-export re-import, resolved generically
rather than per-module).
"""
from __future__ import annotations

import io
from typing import Any

import openpyxl

EXCEPTION_SHEET_HINTS = ("exception", "carry", "discrepan")
VALIDATION_SHEET_HINTS = ("validation", "pass 1", "pass 2", "pass 3")
SUMMARY_SHEET_HINTS = ("summary", "daily", "pivot", "insight")


def _infer_kind(sheet_name: str) -> str:
    lower = sheet_name.lower()
    if any(h in lower for h in EXCEPTION_SHEET_HINTS):
        return "exception"
    if any(h in lower for h in VALIDATION_SHEET_HINTS):
        return "validation"
    if any(h in lower for h in SUMMARY_SHEET_HINTS):
        return "summary"
    return "detail"


def import_report(data: bytes) -> list[dict[str, Any]]:
    """Returns a list of {name, kind, columns, rows} - one per sheet with a
    usable header row. Sheets with no discernible header are skipped."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets: list[dict[str, Any]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header or all(c is None for c in header):
            continue
        columns = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(header)]
        data_rows: list[dict[str, Any]] = []
        for row in rows_iter:
            if row is None or all(c is None for c in row):
                continue
            data_rows.append({columns[i]: (row[i] if i < len(row) else None) for i in range(len(columns))})
        if not data_rows:
            continue
        sheets.append({"name": sheet_name, "kind": _infer_kind(sheet_name), "columns": columns, "rows": data_rows})
    return sheets
