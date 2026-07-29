"""Bank vs Add Money Recon — faithful port of the (bug-fixed)
`Bank_vs_Add_Money_Auto_Reconciliation_Utility_v1_17042026.bas` macro.

Matches Ops "Add Money" credit transactions to bank statement credits by
UTR, extracted from the bank Narrative via prefix-specific rules. See
PRD §7.6 for the full mandatory rule set ported here verbatim.
"""
from __future__ import annotations

import re
from typing import Any

from pydantic import Field

from app.recon.base import (
    InputSlot,
    OptionsBase,
    ParsedData,
    ReconModule,
    ReconOutput,
    ReportSheet,
    RunContext,
    StatCard,
    ValidationFinding,
)
from app.recon.registry import register
from app.recon.utils.currency import round2, to_float
from app.recon.utils.excel_read import read_rows, rows_to_dicts
from app.recon.utils.excel_write import (
    bold_font,
    freeze,
    merge_and_label,
    set_cell,
    set_column_widths,
    solid_fill,
    write_row,
)
from app.recon.utils.matching import KeyedAggregator
from openpyxl import Workbook

CATEGORY_MATCH = {"ADD MONEY", "OPS ADJUSTMENT", "OPS ADJUSTEMENT"}  # typo preserved verbatim from source

COLORS = {
    "MATCHED": ("FFC6EFCE", "FF1E7B1E"),
    "AMOUNT MISMATCH": ("FFFFEB9C", "FF9C5700"),
    "NOT IN BANK": ("FFFFC7CE", "FF9C0006"),
    "NOT IN OPS": ("FFDDCAFF", "FF460082"),
}
HEADER_FILL = "FF1F4E79"


class BankAddMoneyOptions(OptionsBase):
    match_tolerance: float = Field(default=0.01, description="Absolute ₹ tolerance for MATCHED status")


def _extract_numbers(s: str) -> str:
    return "".join(ch for ch in s if ch.isdigit())


def _first_n_digits(s: str, n: int) -> str:
    out: list[str] = []
    for ch in s:
        if ch.isdigit():
            out.append(ch)
            if len(out) == n:
                break
    return "".join(out)


def _extract_between_tilde(s: str) -> str:
    upper = s.upper()
    start = upper.find("~EROUTE~")
    if start < 0:
        return ""
    start += len("~EROUTE~")
    end = upper.find("~", start)
    seg = s[start:] if end < 0 else s[start:end]
    return _extract_numbers(seg.strip())


def _extract_between_slashes(s: str) -> str:
    p1 = s.find("/")
    if p1 < 0:
        return ""
    p2 = s.find("/", p1 + 1)
    seg = s[p1 + 1 :] if p2 < 0 else s[p1 + 1 : p2]
    return _extract_numbers(seg.strip())


def extract_ops_utr(utr_val: str) -> str | None:
    """Returns None when the row should be excluded entirely (TRXV*)."""
    upper = utr_val.strip().upper()
    if upper.startswith("TRXV"):
        return None
    if upper.startswith("RE1"):
        return utr_val.strip()[3:]
    return _extract_numbers(utr_val)


def extract_bank_utr(narrative: str) -> str:
    upper = narrative.strip().upper()
    if upper.startswith("IMPS/"):
        return _first_n_digits(narrative, 12)
    if upper.startswith("IFT/"):
        return _first_n_digits(narrative, 11)
    if upper.startswith("INEFT") or upper.startswith("IRTGS"):
        return _extract_between_tilde(narrative)
    if upper.startswith("NEFT") or upper.startswith("RTGS") or upper.startswith("IMPS"):
        return _extract_between_slashes(narrative)
    if narrative.strip() != "":
        return _extract_numbers(narrative)
    return ""


def _find_col(headers: list[str], name: str) -> str | None:
    lname = name.strip().lower()
    for h in headers:
        if h.strip().lower() == lname:
            return h
    return None


class BankAddMoneyModule(ReconModule):
    key = "bank_addmoney"
    display_name = "Bank vs Add Money Recon"
    description = "Reconciles Ops 'Add Money' credit transactions against bank statement credits by UTR."
    version = "1.0.0"
    input_slots = [
        InputSlot(key="ops", label="Ops File", accept=[".xlsx", ".xls", ".csv"], required=True),
        InputSlot(key="bank", label="Bank Statement", accept=[".xlsx", ".xls", ".csv"], required=True),
    ]
    options_schema = BankAddMoneyOptions

    def parse(self, ctx: RunContext) -> ParsedData:
        ops_file = ctx.files_for("ops")[0]
        bank_file = ctx.files_for("bank")[0]

        ops_rows = read_rows(ops_file.content, ops_file.filename)
        bank_rows = read_rows(bank_file.content, bank_file.filename)
        ops_records = rows_to_dicts(ops_rows)
        bank_records = rows_to_dicts(bank_rows)
        ctx.info(f"Ops File: {len(ops_records)} rows")
        ctx.info(f"Bank Statement: {len(bank_records)} rows")
        return ParsedData(data={"ops": ops_records, "bank": bank_records})

    def validate(self, parsed: ParsedData, ctx: RunContext) -> list[ValidationFinding]:
        findings: list[ValidationFinding] = []
        ops, bank = parsed.data["ops"], parsed.data["bank"]

        if not ops:
            findings.append(
                ValidationFinding(pass_name="Input Integrity", check="Ops File", level="err", message="No rows found")
            )
        else:
            headers = list(ops[0].keys())
            required = ["TransactionCategory", "TransactionType", "UtrNo", "TotalTransactionAmount"]
            missing = [r for r in required if _find_col(headers, r) is None]
            if missing:
                findings.append(
                    ValidationFinding(
                        pass_name="Input Integrity",
                        check="Ops File headers",
                        level="err",
                        message=f"Missing required column(s): {', '.join(missing)}",
                    )
                )
            else:
                findings.append(
                    ValidationFinding(
                        pass_name="Input Integrity", check="Ops File headers", level="ok", message="All required columns present"
                    )
                )

        if not bank:
            findings.append(
                ValidationFinding(pass_name="Input Integrity", check="Bank Statement", level="err", message="No rows found")
            )
        else:
            headers = list(bank[0].keys())
            missing = [r for r in ["Narrative", "Credit"] if _find_col(headers, r) is None]
            if missing:
                findings.append(
                    ValidationFinding(
                        pass_name="Input Integrity",
                        check="Bank Statement headers",
                        level="err",
                        message=f"Missing required column(s): {', '.join(missing)}",
                    )
                )
            else:
                findings.append(
                    ValidationFinding(
                        pass_name="Input Integrity",
                        check="Bank Statement headers",
                        level="ok",
                        message="All required columns present",
                    )
                )
        return findings

    def reconcile(self, parsed: ParsedData, ctx: RunContext, findings: list[ValidationFinding]) -> ReconOutput:
        options = BankAddMoneyOptions(**ctx.options)
        tolerance = options.match_tolerance
        ops_records: list[dict[str, Any]] = parsed.data["ops"]
        bank_records: list[dict[str, Any]] = parsed.data["bank"]

        if any(f.level == "err" for f in findings):
            return ReconOutput(stats=[StatCard(label="Status", value="Failed pre-checks", tone="bad")])

        ops_headers = list(ops_records[0].keys()) if ops_records else []
        bank_headers = list(bank_records[0].keys()) if bank_records else []
        col_cat = _find_col(ops_headers, "TransactionCategory")
        col_type = _find_col(ops_headers, "TransactionType")
        col_utr = _find_col(ops_headers, "UtrNo")
        col_amt = _find_col(ops_headers, "TotalTransactionAmount")
        col_ent = _find_col(ops_headers, "EnterpriseName")
        col_narr = _find_col(bank_headers, "Narrative")
        col_credit = _find_col(bank_headers, "Credit")

        ops_agg = KeyedAggregator()
        ent_by_utr: dict[str, str] = {}
        for i, row in enumerate(ops_records):
            cat = str(row.get(col_cat) or "").strip().upper()
            typ = str(row.get(col_type) or "").strip().upper()
            if cat not in CATEGORY_MATCH or typ != "CREDIT":
                continue
            raw_utr = str(row.get(col_utr) or "").strip()
            utr = extract_ops_utr(raw_utr)
            if utr is None or utr == "":
                continue
            amt = to_float(row.get(col_amt))
            ops_agg.add(utr, amt, i)
            if utr not in ent_by_utr:
                ent_by_utr[utr] = str(row.get(col_ent) or "").strip() if col_ent else ""

        bank_agg = KeyedAggregator()
        for i, row in enumerate(bank_records):
            narrative = str(row.get(col_narr) or "").strip()
            credit = to_float(row.get(col_credit))
            utr = extract_bank_utr(narrative)
            if utr and credit != 0:
                bank_agg.add(utr, credit, i)

        detail_rows: list[dict[str, Any]] = []
        matched_bank_keys: set[str] = set()
        n_matched = n_mismatch = n_not_in_bank = 0
        matched_ops_total = matched_bank_total = 0.0

        for utr in ops_agg.keys():
            ops_amt = round2(ops_agg.get(utr) or 0.0)
            ent_name = ent_by_utr.get(utr, "")
            if bank_agg.exists(utr):
                bank_amt = round2(bank_agg.get(utr) or 0.0)
                difference = round2(ops_amt - bank_amt)
                status = "MATCHED" if abs(difference) < tolerance else "AMOUNT MISMATCH"
                matched_bank_keys.add(utr)
                if status == "MATCHED":
                    n_matched += 1
                    matched_ops_total += ops_amt
                    matched_bank_total += bank_amt
                else:
                    n_mismatch += 1
            else:
                bank_amt = None
                difference = ops_amt
                status = "NOT IN BANK"
                n_not_in_bank += 1
            detail_rows.append(
                {
                    "Final UTR": utr,
                    "Enterprise Name": ent_name,
                    "Ops Amount": ops_amt,
                    "Bank Amount": bank_amt,
                    "Difference": round2(difference),
                    "Status": status,
                }
            )

        n_not_in_ops = 0
        for utr in bank_agg.keys():
            if utr in matched_bank_keys:
                continue
            bank_amt = round2(bank_agg.get(utr) or 0.0)
            detail_rows.append(
                {
                    "Final UTR": utr,
                    "Enterprise Name": "",
                    "Ops Amount": None,
                    "Bank Amount": bank_amt,
                    "Difference": round2(-bank_amt),
                    "Status": "NOT IN OPS",
                }
            )
            n_not_in_ops += 1

        net_unreconciled = round2(sum(r["Difference"] for r in detail_rows if r["Status"] == "AMOUNT MISMATCH"))

        columns = ["Final UTR", "Enterprise Name", "Ops Amount", "Bank Amount", "Difference", "Status"]
        exception_rows = [r for r in detail_rows if r["Status"] != "MATCHED"]

        stats = [
            StatCard(label="Total Ops UTRs", value=str(len(ops_agg.keys())), tone="neutral"),
            StatCard(label="Matched", value=str(n_matched), tone="good"),
            StatCard(label="Amount Mismatch", value=str(n_mismatch), tone="bad" if n_mismatch else "good"),
            StatCard(label="Not in Bank", value=str(n_not_in_bank), tone="warn" if n_not_in_bank else "good"),
            StatCard(label="Not in Ops", value=str(n_not_in_ops), tone="warn" if n_not_in_ops else "good"),
            StatCard(label="Total Ops Amount (Matched)", value=f"₹{matched_ops_total:,.2f}"),
            StatCard(label="Total Bank Amount (Matched)", value=f"₹{matched_bank_total:,.2f}"),
            StatCard(
                label="Net Unreconciled Difference",
                value=f"₹{net_unreconciled:,.2f}",
                tone="bad" if abs(net_unreconciled) > 0 else "good",
            ),
        ]

        sheets = [
            ReportSheet(name="Reconciliation", kind="detail", columns=columns, rows=detail_rows),
            ReportSheet(name="Exceptions", kind="exception", columns=columns, rows=exception_rows),
        ]
        return ReconOutput(stats=stats, sheets=sheets, findings=[], raw={"detail_rows": detail_rows})

    def build_report(self, output: ReconOutput, ctx: RunContext) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Reconciliation"
        columns = ["Final UTR", "Enterprise Name", "Ops Amount", "Bank Amount", "Difference", "Status"]
        set_column_widths(ws, {1: 22, 2: 26, 3: 16, 4: 16, 5: 14, 6: 20})

        write_row(
            ws, 1, columns, cell_font=None, fill=solid_fill(HEADER_FILL), align=None
        )
        for i, name in enumerate(columns, start=1):
            set_cell(ws, 1, i, name, cell_font=None, fill=solid_fill(HEADER_FILL))
            bold_font(ws.cell(row=1, column=i), color="FFFFFFFF")
        freeze(ws, "A2")

        detail_rows = output.raw["detail_rows"] if output.raw else []
        row_idx = 2
        for row in detail_rows:
            bg, fg = COLORS.get(row["Status"], ("FFFFFFFF", "FF000000"))
            values = [row["Final UTR"], row["Enterprise Name"], row["Ops Amount"], row["Bank Amount"], row["Difference"], row["Status"]]
            write_row(ws, row_idx, values, fill=solid_fill(bg))
            ws.cell(row=row_idx, column=3).number_format = "#,##0.00"
            ws.cell(row=row_idx, column=4).number_format = "#,##0.00"
            ws.cell(row=row_idx, column=5).number_format = "#,##0.00"
            bold_font(ws.cell(row=row_idx, column=6), color=fg)
            row_idx += 1

        data_end = row_idx - 1
        summary_row = data_end + 3
        merge_and_label(
            ws, summary_row, 1, summary_row, 6, "RECONCILIATION SUMMARY", fill=solid_fill(HEADER_FILL)
        )
        bold_font(ws.cell(row=summary_row, column=1), color="FFFFFFFF")
        summary_defs = [
            ("Matched", "MATCHED"),
            ("Amount Mismatch", "AMOUNT MISMATCH"),
            ("Not in Bank", "NOT IN BANK"),
            ("Not in Ops", "NOT IN OPS"),
        ]
        for offset, (label, status_key) in enumerate(summary_defs):
            r = summary_row + 1 + offset
            bg, fg = COLORS[status_key]
            set_cell(ws, r, 1, label, fill=solid_fill(bg))
            bold_font(ws.cell(row=r, column=1), color=fg)
            set_cell(ws, r, 2, f'=COUNTIF(F2:F{data_end},"{status_key}")')

        totals_row = summary_row + 6
        set_cell(ws, totals_row, 1, "Total Ops Amount (Matched)")
        set_cell(ws, totals_row, 2, f'=SUMIF(F2:F{data_end},"MATCHED",C2:C{data_end})', number_format="#,##0.00")
        set_cell(ws, totals_row + 1, 1, "Total Bank Amount (Matched)")
        set_cell(ws, totals_row + 1, 2, f'=SUMIF(F2:F{data_end},"MATCHED",D2:D{data_end})', number_format="#,##0.00")
        set_cell(ws, totals_row + 2, 1, "Net Unreconciled Difference")
        set_cell(
            ws, totals_row + 2, 2, f'=SUMIF(F2:F{data_end},"AMOUNT MISMATCH",E2:E{data_end})', number_format="#,##0.00"
        )
        for r in range(totals_row, totals_row + 3):
            bold_font(ws.cell(row=r, column=1))
        import io

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


register(BankAddMoneyModule())
