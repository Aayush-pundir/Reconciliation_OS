"""NETC FASTag Recon — faithful port of Fastag_Recon_Tool_v2 (one DSR file
= one settlement block/cycle; two-pass bank matching incl. amount-based
narrative-mismatch reclassification). See PRD §7.4.
"""
from __future__ import annotations

from datetime import date, datetime
from typing import Any

from pydantic import Field

from app.recon.base import (
    InputSlot,
    OptionsBase,
    ParsedData,
    ReconModule,
    ReconOutput,
    ReportSheet,
    RunContext,
    StatCard,
    ValidationFinding,
)
from app.recon.registry import register
from app.recon.utils.currency import round2, to_float, to_int
from app.recon.utils.excel_read import read_rows
from app.recon.utils.excel_write import bold_font, freeze, set_cell, set_column_widths, solid_fill, write_row
from openpyxl import Workbook


class FastagOptions(OptionsBase):
    bridge_tolerance: float = Field(default=0.05, description="₹ tolerance for the overall recon bridge check")


def _get(row: list[Any], idx: int) -> Any:
    return row[idx] if row is not None and idx < len(row) else None


def fmt_date(v: Any) -> str:
    if isinstance(v, (datetime, date)):
        return f"{v.day:02d}-{v.month:02d}-{v.year:04d}"
    if v is None:
        return ""
    return str(v).strip()[:10]


def parse_single_dsr(rows: list[list[Any]], filename: str, cycle_counter: dict[str, int]) -> dict[str, Any] | None:
    header_idx = -1
    for i in range(min(len(rows), 5)):
        if _get(rows[i], 0) == "Settlement Date":
            header_idx = i
            break
    if header_idx == -1:
        return None

    block_date = None
    for i in range(header_idx + 1, len(rows)):
        v = _get(rows[i], 0)
        if v is not None and v != "Settlement Date" and str(v).strip():
            block_date = fmt_date(v)
            break
    if not block_date:
        return None

    cycle_counter[block_date] = cycle_counter.get(block_date, 0) + 1
    cycle_num = cycle_counter[block_date]

    grand_total: list[Any] | None = None
    gst_amt = 0.0
    txn_count = 0
    txn_amt = 0.0

    for k in range(header_idx + 1, len(rows)):
        r = rows[k]
        if not r:
            continue
        if str(_get(r, 5) or "").strip() == "INWARD GST":
            gst_amt = to_float(_get(r, 18))
        if _get(r, 1) == "Total" and _get(r, 4) == "Total":
            grand_total = r
        if _get(r, 8) == "DEBIT" and _get(r, 9) in ("TOLL", "PARKING"):
            txn_count += to_int(_get(r, 10))
            txn_amt += to_float(_get(r, 12))

    if grand_total is None:
        return None

    svc_total = to_float(_get(grand_total, 18))
    taxable = round2(svc_total - gst_amt)
    gst = round2(gst_amt)
    net = round2(to_float(_get(grand_total, 27)))

    return {
        "date": block_date,
        "cycle": cycle_num,
        "filename": filename,
        "cnt": txn_count,
        "txnAmt": round2(txn_amt),
        "taxable": taxable,
        "gst": gst,
        "svcTotal": round2(svc_total),
        "net": net,
    }


def parse_bank(rows: list[list[Any]]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for i in range(1, len(rows)):
        row = rows[i]
        if not row:
            continue
        narr = str(_get(row, 1) or "")
        if "_00" not in narr:
            continue
        parts = narr.split("_00")
        if len(parts) < 2 or len(parts[1]) < 10:
            continue
        code = parts[1]
        date_str = code[:8]
        try:
            cycle = int(code[8:])
        except ValueError:
            continue
        dd = f"{date_str[6:8]}-{date_str[4:6]}-{date_str[0:4]}"
        pay_date = fmt_date(_get(row, 0))
        debit = to_float(_get(row, 5))
        entries.append(
            {"key": f"{dd}|{cycle}", "dsrDate": dd, "cycle": cycle, "payDate": pay_date, "debit": debit, "narr": narr}
        )
    return entries


class FastagModule(ReconModule):
    key = "fastag"
    display_name = "NETC FASTag Recon"
    description = "Reconciles per-cycle FASTag DSR settlement files against a bank statement."
    version = "1.0.0"
    input_slots = [
        InputSlot(
            key="dsr",
            label="DSR Files",
            accept=[".xlsx", ".xls"],
            multiple=True,
            required=True,
            help_text="One file per settlement cycle.",
        ),
        InputSlot(key="bank", label="Bank Statement", accept=[".xlsx", ".xls"], required=True),
    ]
    options_schema = FastagOptions

    def parse(self, ctx: RunContext) -> ParsedData:
        dsr_files = sorted(ctx.files_for("dsr"), key=lambda f: f.filename)
        bank_file = ctx.files_for("bank")[0]

        blocks: list[dict[str, Any]] = []
        skipped: list[dict[str, str]] = []
        cycle_counter: dict[str, int] = {}
        for f in dsr_files:
            try:
                rows = read_rows(f.content, f.filename)
                block = parse_single_dsr(rows, f.filename, cycle_counter)
                if block:
                    blocks.append(block)
                else:
                    skipped.append({"file": f.filename, "reason": "not a recognizable DSR file"})
            except Exception as exc:  # noqa: BLE001
                skipped.append({"file": f.filename, "reason": str(exc)})

        bank_rows = read_rows(bank_file.content, bank_file.filename)
        bank_entries = parse_bank(bank_rows)

        ctx.info(f"Parsed {len(blocks)} DSR cycle(s) from {len(dsr_files)} file(s), {len(bank_entries)} bank entries")
        return ParsedData(data={"blocks": blocks, "bank_entries": bank_entries, "skipped": skipped, "file_count": len(dsr_files)})

    def validate(self, parsed: ParsedData, ctx: RunContext) -> list[ValidationFinding]:
        findings: list[ValidationFinding] = []
        if not parsed.data["blocks"]:
            findings.append(
                ValidationFinding(pass_name="Input Integrity", check="DSR blocks", level="err", message="No DSR blocks parsed")
            )
        else:
            findings.append(
                ValidationFinding(
                    pass_name="Input Integrity", check="DSR blocks", level="ok", message=f"{len(parsed.data['blocks'])} block(s) parsed"
                )
            )
        if not parsed.data["bank_entries"]:
            findings.append(
                ValidationFinding(pass_name="Input Integrity", check="Bank entries", level="err", message="No bank entries found")
            )
        if parsed.data["skipped"]:
            findings.append(
                ValidationFinding(
                    pass_name="Input Integrity",
                    check="Skipped DSR files",
                    level="warn",
                    message=f"{len(parsed.data['skipped'])} file(s) skipped",
                )
            )
        return findings

    def reconcile(self, parsed: ParsedData, ctx: RunContext, findings: list[ValidationFinding]) -> ReconOutput:
        if any(f.level == "err" for f in findings):
            return ReconOutput(stats=[StatCard(label="Status", value="Failed pre-checks", tone="bad")])

        options = FastagOptions(**ctx.options)
        blocks: list[dict[str, Any]] = parsed.data["blocks"]
        bank_entries: list[dict[str, Any]] = parsed.data["bank_entries"]

        bank_dict = {b["key"]: b for b in bank_entries}  # last-write-wins, matches source
        matched_keys = {f"{blk['date']}|{blk['cycle']}" for blk in blocks if f"{blk['date']}|{blk['cycle']}" in bank_dict}

        bank_used = set(matched_keys)
        reclass_map: dict[str, dict[str, Any]] = {}
        for blk in blocks:
            key = f"{blk['date']}|{blk['cycle']}"
            if key in matched_keys or blk["net"] == 0:
                continue
            net_abs = abs(blk["net"])
            for b in bank_entries:
                if b["cycle"] == blk["cycle"] and b["key"] not in bank_used and abs(round2(b["debit"] - net_abs)) < 0.01:
                    reclass_map[key] = {**b, "reclassNote": f"Bank narrative {b['narr']} → DSR date {blk['date']}"}
                    bank_used.add(b["key"])
                    break

        dates = sorted({blk["date"] for blk in blocks})
        last_date = dates[-1] if dates else None

        recon_rows: list[dict[str, Any]] = []
        for blk in blocks:
            key = f"{blk['date']}|{blk['cycle']}"
            net_abs = abs(blk["net"] or 0)
            bank_entry = None
            diff = None
            if key in bank_dict:
                bank_entry = bank_dict[key]
                diff = round2(bank_entry["debit"] - net_abs)
                if blk["net"] == 0:
                    status = "Zero DSR – No Debit"
                elif abs(diff) < 0.01:
                    status = "Matched"
                else:
                    status = f"Amt Diff ₹{diff:.2f}"
            elif key in reclass_map:
                bank_entry = reclass_map[key]
                diff = round2(bank_entry["debit"] - net_abs)
                status = "Matched – Narrative Mismatch*"
            elif blk["net"] == 0:
                status = "Zero DSR – No Debit"
            elif blk["date"] == last_date:
                status = "Pending – Next Month Debit"
            else:
                status = "Unmatched – Investigate"

            recon_rows.append(
                {
                    "DSR Date": blk["date"],
                    "Cycle": blk["cycle"],
                    "File": blk["filename"],
                    "Txn Count": blk["cnt"],
                    "Txn Amt Dr": blk["txnAmt"],
                    "Taxable Value": blk["taxable"],
                    "GST @18%": blk["gst"],
                    "Total Svc Fee": blk["svcTotal"],
                    "Final Net Amt": blk["net"],
                    "Bank Debit": bank_entry["debit"] if bank_entry else None,
                    "Bank Pay Date": bank_entry["payDate"] if bank_entry else None,
                    "Difference": diff,
                    "Status": status,
                }
            )

        matched = sum(1 for r in recon_rows if r["Status"] == "Matched")
        reclass = sum(1 for r in recon_rows if "Narrative" in r["Status"])
        pending = sum(1 for r in recon_rows if "Pending" in r["Status"])
        zero = sum(1 for r in recon_rows if "Zero" in r["Status"])
        unmatched = sum(1 for r in recon_rows if "Unmatched" in r["Status"])

        total_txn_amt = round2(sum(r["Txn Amt Dr"] for r in recon_rows))
        total_taxable = round2(sum(r["Taxable Value"] for r in recon_rows))
        total_gst = round2(sum(r["GST @18%"] for r in recon_rows))
        total_svc = round2(sum(r["Total Svc Fee"] for r in recon_rows))
        total_net_abs = round2(sum(abs(r["Final Net Amt"] or 0) for r in recon_rows))
        bank_total = round2(sum(b["debit"] for b in bank_entries))
        pending_amt = round2(sum(abs(r["Final Net Amt"] or 0) for r in recon_rows if "Pending" in r["Status"]))
        bridge_diff = round2(bank_total + pending_amt - total_net_abs)

        stats = [
            StatCard(label="DSR Cycles Loaded", value=str(len(recon_rows)), sub=f"From {parsed.data['file_count']} files"),
            StatCard(label="Matched", value=str(matched + reclass), tone="good", sub=f"Exact: {matched} | Narrative mismatch: {reclass}"),
            StatCard(label="Pending / Exceptions", value=str(pending + unmatched), tone="warn" if pending + unmatched else "good"),
            StatCard(
                label="Net Difference",
                value="₹0.00 ✓" if abs(bridge_diff) < options.bridge_tolerance else f"₹{abs(bridge_diff):,.2f}",
                tone="good" if abs(bridge_diff) < options.bridge_tolerance else "bad",
                sub=f"Bank ₹{bank_total:,.2f} + Pending ₹{pending_amt:,.2f}",
            ),
            StatCard(label="Total Svc Fee (incl. GST)", value=f"₹{total_svc:,.2f}"),
            StatCard(label="DSR Net Payable", value=f"₹{total_net_abs:,.2f}"),
        ]

        findings_out: list[ValidationFinding] = []
        if abs(bridge_diff) < options.bridge_tolerance:
            findings_out.append(
                ValidationFinding(
                    pass_name="Key Findings & Actions",
                    check="Recon bridge",
                    level="ok",
                    message=f"FULL RECON ACHIEVED — DSR ₹{total_net_abs:,.2f} = Bank ₹{bank_total:,.2f} + Pending ₹{pending_amt:,.2f}.",
                )
            )
        else:
            findings_out.append(
                ValidationFinding(pass_name="Key Findings & Actions", check="Recon bridge", level="err", message=f"DIFFERENCE: ₹{bridge_diff:,.2f} — investigate before closing.")
            )
        if reclass:
            amt = round2(sum(abs(r["Final Net Amt"] or 0) for r in recon_rows if "Narrative" in r["Status"]))
            findings_out.append(
                ValidationFinding(
                    pass_name="Key Findings & Actions",
                    check="Narrative mismatch",
                    level="warn",
                    message=f"{reclass} entries, ₹{amt:,.2f} — amounts correct; bank narrative carries wrong DSR date. Flag to NPCI/bank.",
                )
            )
        if pending:
            findings_out.append(
                ValidationFinding(
                    pass_name="Key Findings & Actions",
                    check="Pending next month",
                    level="warn",
                    message=f"{pending} entries, ₹{pending_amt:,.2f} — verify in next bank statement.",
                )
            )
        if unmatched:
            findings_out.append(
                ValidationFinding(
                    pass_name="Key Findings & Actions",
                    check="Unmatched DSR",
                    level="err",
                    message=f"{unmatched} entries with no bank debit found — escalate immediately.",
                )
            )

        columns = list(recon_rows[0].keys()) if recon_rows else []
        exception_rows = [r for r in recon_rows if r["Status"] != "Matched"]

        sheets = [
            ReportSheet(name="Recon Detail", kind="detail", columns=columns, rows=recon_rows),
            ReportSheet(name="Exceptions", kind="exception", columns=columns, rows=exception_rows),
        ]
        return ReconOutput(stats=stats, sheets=sheets, findings=findings_out, raw={"columns": columns, "rows": recon_rows})

    def build_report(self, output: ReconOutput, ctx: RunContext) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Recon Detail"
        columns = output.raw["columns"]
        set_column_widths(ws, {i: 14 for i in range(1, len(columns) + 1)} | {3: 26, 13: 22})
        for i, name in enumerate(columns, start=1):
            set_cell(ws, 1, i, name, fill=solid_fill("FF1F4E79"))
            bold_font(ws.cell(row=1, column=i), color="FFFFFFFF")
        freeze(ws, "A2")
        row_idx = 2
        for row in output.raw["rows"]:
            values = [row.get(c) for c in columns]
            write_row(ws, row_idx, values)
            row_idx += 1

        summary = wb.create_sheet("Summary")
        for i, stat in enumerate(output.stats, start=1):
            set_cell(summary, i, 1, stat.label)
            set_cell(summary, i, 2, stat.value)
            set_cell(summary, i, 3, stat.sub or "")
        set_column_widths(summary, {1: 28, 2: 22, 3: 40})

        exceptions = wb.create_sheet("Exceptions")
        for i, name in enumerate(columns, start=1):
            set_cell(exceptions, 1, i, name, fill=solid_fill("FFB71C1C"))
            bold_font(exceptions.cell(row=1, column=i), color="FFFFFFFF")
        row_idx = 2
        for row in output.raw["rows"]:
            if row["Status"] == "Matched":
                continue
            write_row(exceptions, row_idx, [row.get(c) for c in columns])
            row_idx += 1
        set_column_widths(exceptions, {i: 14 for i in range(1, len(columns) + 1)})

        import io

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


register(FastagModule())
