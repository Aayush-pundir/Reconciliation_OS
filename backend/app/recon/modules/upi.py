"""UPI NTSL Recon — faithful port of UPI_Recon_Utility (12-cycle NTSL
parsing, remarks classifier, 5-pattern bank narrative parser, 3-pass
validation). See PRD §7.3.

Pass 1 (input integrity) and Pass 2 (data integrity) run in `validate()`
since they only need parsed NTSL/bank data. Pass 3 (reconciliation
integrity) needs match results, which only exist after `reconcile()`
computes them - so its findings are appended to `ReconOutput.findings`
from inside `reconcile()`, exactly like RuPay's and FASTag's reconcile-time
findings. The engine persists both sets identically either way.
"""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from typing import Any

from pydantic import Field

from app.recon.base import (
    InputSlot,
    OptionsBase,
    ParsedData,
    ReconModule,
    ReconOutput,
    ReportSheet,
    RunContext,
    StatCard,
    ValidationFinding,
)
from app.recon.registry import register
from app.recon.utils.currency import round2, round4, to_float
from app.recon.utils.excel_read import read_rows
from app.recon.utils.excel_write import bold_font, freeze, set_cell, set_column_widths, solid_fill, write_row
from openpyxl import Workbook

CYCLE_ORDER = ["10C", "1C", "2C", "3C", "4C", "5C", "6C", "7C", "8C", "9C", "DC1", "DC2"]
T_CYCLES = {"3C", "4C", "5C", "6C", "7C", "8C", "9C"}
T1_CYCLES = {"10C", "1C", "2C", "DC1", "DC2"}
CYCLE_MAP = {f"{i:02d}": (f"{i}C" if i <= 9 else "10C") for i in range(1, 11)}
REMARKS_ORDER = [
    "Transaction Amount", "Switching Fee", "Switching Fee Gst", "Approved Fee", "Approved Fee Gst",
    "Net Adjusted Amount", "Penalty", "GST on Penalty", "TDS on GST",
]
COMPANY = "EROUTE TECHNOLOGIES PVT LTD"
MATCH_TOL = 1.0
SKIP_RE = re.compile(r"^\s*(total|grand|sub.?total|net settlement|settlement amount|opening|closing)", re.IGNORECASE)


def _get(row: list[Any], idx: int) -> Any:
    return row[idx] if row is not None and idx < len(row) else None


def date_key(d: date) -> str:
    return d.strftime("%Y-%m-%d")


def classify_remarks(desc: str) -> str | None:
    d = desc or ""
    if re.search(r"TDS", d, re.IGNORECASE):
        return "TDS on GST"
    if re.search(r"gst\s+on", d, re.IGNORECASE):
        return "GST on Penalty"
    if re.search(r"penalty", d, re.IGNORECASE):
        return "Penalty"
    if re.search(r"Net Adjusted Amount", d, re.IGNORECASE):
        return "Net Adjusted Amount"
    if re.search(r"Transaction Amount", d, re.IGNORECASE):
        return "Transaction Amount"
    if re.search(r"Switching Fee\s*Gst|Switching Fee\s*GST", d, re.IGNORECASE):
        return "Switching Fee Gst"
    if re.search(r"Switching Fee", d, re.IGNORECASE):
        return "Switching Fee"
    if re.search(r"Fee\s*Gst|Fee\s*GST", d, re.IGNORECASE):
        return "Approved Fee Gst"
    if re.search(r"Fee", d, re.IGNORECASE):
        return "Approved Fee"
    return None


def parse_ntsl(rows: list[list[Any]], filename: str) -> list[dict[str, Any]]:
    if len(rows) < 5:
        raise ValueError("File too short")

    header_text = str(_get(rows[2], 0) or "")
    settle_date: date | None = None
    cycle = ""

    dm = re.search(r"as on (\d{2})-(\d{2})-(\d{4})", header_text, re.IGNORECASE)
    cm = re.search(r"\(\s*([A-Z0-9]+)\s+\d{2}:\d{2}:\d{2}", header_text, re.IGNORECASE)
    if dm and cm:
        settle_date = date(int(dm.group(3)), int(dm.group(2)), int(dm.group(1)))
        cycle = cm.group(1).upper()
    else:
        fm = re.search(r"NTSLERT(\d{2})(\d{2})(\d{2})_(\w+)\.(xls|xlsx|xlsm|xlsb)$", filename, re.IGNORECASE)
        if not fm:
            raise ValueError("Cannot parse date/cycle from header or filename")
        settle_date = date(2000 + int(fm.group(3)), int(fm.group(2)), int(fm.group(1)))
        cycle = fm.group(4).upper()

    cycle = cycle.replace("DC01", "DC1").replace("DC02", "DC2")
    is_dc = cycle in ("DC1", "DC2")

    header_row = [str(c or "").lower() for c in (rows[4] or [])]

    def find_idx(*keywords: str) -> int:
        for kw in keywords:
            for i, h in enumerate(header_row):
                if kw in h:
                    return i
        return -1

    col_desc = find_idx("description", "particular", "remark")
    col_desc = 0 if col_desc < 0 else col_desc
    col_txn = find_idx("no. of", "no of", "count", "txn")
    col_txn = 1 if col_txn < 0 else col_txn
    col_debit = find_idx("debit")
    col_debit = 2 if col_debit < 0 else col_debit
    col_credit = find_idx("credit")
    col_credit = 3 if col_credit < 0 else col_credit

    out: list[dict[str, Any]] = []
    for r in rows[5:]:
        if not r or _get(r, col_desc) in (None, ""):
            continue
        desc = str(_get(r, col_desc)).strip()
        if not desc or SKIP_RE.match(desc):
            continue
        remark = classify_remarks(desc)
        if not remark:
            continue
        debit = to_float(_get(r, col_debit))
        credit = to_float(_get(r, col_credit))
        txn = round(to_float(_get(r, col_txn)))
        if debit == 0 and credit == 0:
            continue
        if is_dc and remark not in ("Net Adjusted Amount", "Penalty", "GST on Penalty", "TDS on GST"):
            continue
        out.append({"date": settle_date, "cycle": cycle, "desc": desc, "noTxns": txn, "debit": debit, "credit": credit, "remarks": remark})
    return out


def parse_bank_narrative(raw: str) -> tuple[date, str] | None:
    n = (raw or "").strip()

    m = re.search(r"Eroute UPI Setl\s+(\S+)\s+(\d{6})", n, re.IGNORECASE)
    if m:
        cyc = m.group(1).upper()
        s = m.group(2)
        try:
            return date(2000 + int(s[4:6]), int(s[2:4]), int(s[0:2])), cyc
        except ValueError:
            pass

    m = re.search(r"DC0([12])\s+UPI EROUTE TECH PVT LTD_NPCI", n, re.IGNORECASE)
    if m:
        dm = re.search(r"NPCI00(\d{4})(\d{2})(\d{2})", n, re.IGNORECASE)
        if dm:
            try:
                return date(int(dm.group(1)), int(dm.group(2)), int(dm.group(3))), f"DC{m.group(1)}"
            except ValueError:
                pass

    m = re.search(r"NPCI00(\d{4})(\d{2})(\d{2})(\d{2})$", n, re.IGNORECASE)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))), CYCLE_MAP.get(m.group(4), m.group(4))
        except ValueError:
            pass

    m_dc = re.search(r"DC([12])", n, re.IGNORECASE)
    m_dt = re.search(r"NPCI00(\d{4})(\d{2})(\d{2})", n, re.IGNORECASE)
    if m_dc and m_dt:
        try:
            return date(int(m_dt.group(1)), int(m_dt.group(2)), int(m_dt.group(3))), f"DC{m_dc.group(1)}"
        except ValueError:
            pass

    m = re.search(r"NPCI00(\d{4})(\d{4})(\d{2})$", n, re.IGNORECASE)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)[0:2]), int(m.group(2)[2:4])), CYCLE_MAP.get(m.group(3), m.group(3))
        except ValueError:
            pass

    return None


def load_bank(rows: list[list[Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for i, r in enumerate(rows[1:], start=1):
        if not r or _get(r, 0) is None:
            continue
        narrative = str(_get(r, 0)).strip()
        debit = to_float(_get(r, 4))
        credit = to_float(_get(r, 5))
        parsed = parse_bank_narrative(narrative)
        out.append(
            {
                "rowNum": i + 1, "narrative": narrative,
                "erouteRem": str(_get(r, 1) or ""), "custRef": str(_get(r, 2) or ""),
                "bankDebit": debit, "bankCredit": credit,
                "ntslDate": parsed[0] if parsed else None, "ntslCycle": parsed[1] if parsed else None,
            }
        )
    return out


def _run_pass1(ntsl_files: list[Any], bank_present: bool) -> list[ValidationFinding]:
    items: list[ValidationFinding] = []
    if not ntsl_files:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="File count", level="err", message="No NTSL files found"))
    elif len(ntsl_files) > 500:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="File count", level="warn", message=f"{len(ntsl_files)} files — exceeds expected max of 500"))
    else:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="File count", level="ok", message=f"{len(ntsl_files)} NTSL files found"))

    bad = [f for f in ntsl_files if not re.search(r"UPI_NTSLERT\d{6}_\w+\.(xls|xlsx|xlsm|xlsb)$", f.filename, re.IGNORECASE)]
    if bad:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Filename pattern", level="warn", message=f"{len(bad)} file(s) have unexpected naming pattern"))
    else:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Filename pattern", level="ok", message="All filenames match UPI_NTSLERT{DDMMYY}_{CYCLE}.xls"))

    coverage: dict[str, set[str]] = {}
    for f in ntsl_files:
        m = re.search(r"NTSLERT(\d{6})_(\w+)\.(xls|xlsx|xlsm|xlsb)$", f.filename, re.IGNORECASE)
        if m:
            dk = m.group(1)
            cyc = m.group(2).upper().replace("DC01", "DC1").replace("DC02", "DC2")
            coverage.setdefault(dk, set()).add(cyc)
    incomplete = [dk for dk, cycles in coverage.items() if not set(CYCLE_ORDER).issubset(cycles)]
    if incomplete:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Cycle coverage", level="warn", message=f"{len(incomplete)} date(s) have incomplete cycle coverage"))
    else:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Cycle coverage", level="ok", message=f"All {len(coverage)} date(s) have full 12-cycle coverage"))

    seen: set[str] = set()
    dupes = 0
    for f in ntsl_files:
        k = f.filename.upper()
        if k in seen:
            dupes += 1
        seen.add(k)
    if dupes:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Duplicates", level="warn", message=f"{dupes} duplicate filename(s) detected"))
    else:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Duplicates", level="ok", message="No duplicate NTSL filenames detected"))

    if not bank_present:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Bank statement", level="err", message="Bank statement not selected"))
    else:
        items.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check="Bank statement", level="ok", message="Bank statement present"))
    return items


def _run_pass2(all_rows: list[dict[str, Any]], ntsl_net: dict[str, float], file_count: int) -> list[ValidationFinding]:
    items: list[ValidationFinding] = []
    if not all_rows:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Row count", level="err", message="No rows parsed from NTSL files"))
        return items
    items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Row count", level="ok", message=f"{len(all_rows)} working rows parsed from {file_count} NTSL files"))

    zero = sum(1 for r in all_rows if r["debit"] == 0 and r["credit"] == 0)
    if zero:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Zero-value rows", level="err", message=f"{zero} zero-debit-zero-credit rows found"))
    else:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Zero-value rows", level="ok", message="No zero-value rows in working data"))

    recompute: dict[str, dict[str, float]] = {}
    for r in all_rows:
        k = date_key(r["date"]) + "_" + r["cycle"]
        acc = recompute.setdefault(k, {"d": 0.0, "c": 0.0})
        acc["d"] += r["debit"]
        acc["c"] += r["credit"]
    mismatches = 0
    for k, v in recompute.items():
        expected = round4(v["c"] - v["d"])
        stored = round4(ntsl_net.get(k, 0.0))
        if abs(expected - stored) > 0.02:
            mismatches += 1
    if mismatches:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Cycle net re-computation", level="err", message=f"{mismatches} cycle(s) have net inconsistency on re-computation"))
    else:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Cycle net re-computation", level="ok", message=f"All {len(recompute)} cycle nets verified — Credit−Debit consistent"))

    total_d = sum(r["debit"] for r in all_rows)
    total_c = sum(r["credit"] for r in all_rows)
    net_rows = total_c - total_d
    net_cyc = sum(ntsl_net.values())
    gap = abs(net_rows - net_cyc)
    if gap > 1:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Grand total cross-check", level="err", message=f"Grand total mismatch: row-level vs cycle-sum diff = ₹{gap:.2f}"))
    else:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Grand total cross-check", level="ok", message=f"PASSED — Net = ₹{net_cyc:,.2f} (diff = {gap:.4f})"))

    no_remark = sum(1 for r in all_rows if not r["remarks"])
    if no_remark:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Remarks classification", level="warn", message=f"{no_remark} rows have no Remarks classification"))
    else:
        items.append(ValidationFinding(pass_name="Pass 2 — Data Integrity", check="Remarks classification", level="ok", message="All rows have valid Remarks classification"))
    return items


def _run_pass3(bank_data: list[dict[str, Any]], ntsl_net: dict[str, float], match_res: list[dict[str, Any]]) -> tuple[list[ValidationFinding], list[dict[str, Any]]]:
    items: list[ValidationFinding] = []

    parse_fail = sum(1 for b in bank_data if not b["ntslDate"])
    if parse_fail:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Narrative parsing", level="err", message=f"{parse_fail} bank entries failed narrative parsing"))
    else:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Narrative parsing", level="ok", message=f"All {len(bank_data)} bank narratives parsed successfully"))

    remap: dict[str, float] = {}
    for b in bank_data:
        if not b["ntslDate"]:
            continue
        k = date_key(b["ntslDate"]) + "_" + b["ntslCycle"]
        nv = ntsl_net.get(k)
        if nv is not None:
            remap[k] = round4((b["bankCredit"] - b["bankDebit"]) - nv)
    verify_mismatches = 0
    for r in match_res:
        if not r["ntslDate"] or "Prior Month" in (r["matchStatus"] or ""):
            continue
        k = date_key(r["ntslDate"]) + "_" + r["ntslCycle"]
        stored = round4(r["diff"] or 0)
        if k in remap and abs(stored - remap[k]) > 0.02:
            verify_mismatches += 1
    if verify_mismatches:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Re-match verification", level="err", message=f"{verify_mismatches} match result(s) differ on re-validation"))
    else:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Re-match verification", level="ok", message="PASSED — all results confirmed on second pass"))

    apr_mr = [r for r in match_res if "Prior" not in (r["matchStatus"] or "")]
    act_disc = [r for r in apr_mr if r["diff"] is not None and abs(r["diff"]) >= MATCH_TOL]
    if act_disc:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Discrepancies", level="warn", message=f"{len(act_disc)} cycle(s) have Bank vs NTSL difference ≥ ₹{MATCH_TOL} — real reconciling items"))
    else:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Discrepancies", level="ok", message="No Bank vs NTSL discrepancies found — all within ₹1 tolerance"))

    apr_b_net = sum(r["bankCredit"] - r["bankDebit"] for r in apr_mr)
    apr_n_net = sum((r["ntslNet"] or 0) for r in apr_mr)
    gap = round2(apr_b_net - apr_n_net)
    if abs(gap) > 1:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Grand total match", level="warn", message=f"Grand total difference: Bank net vs NTSL net = ₹{gap:,.2f}"))
    else:
        items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Grand total match", level="ok", message=f"PASSED — Bank net ≈ NTSL net (diff = ₹{gap:.2f})"))

    matched = sum(1 for r in apr_mr if r["matchStatus"] == "Matched")
    prior = sum(1 for r in match_res if "Prior" in (r["matchStatus"] or ""))
    items.append(ValidationFinding(pass_name="Pass 3 — Reconciliation Integrity", check="Match rate", level="ok", message=f"Match rate: {matched}/{len(apr_mr)} cycles matched | Prior month: {prior}"))
    return items, act_disc


class UPIOptions(OptionsBase):
    pass


class UPIModule(ReconModule):
    key = "upi"
    display_name = "UPI NTSL Recon"
    description = "Reconciles UPI NPCI NTSL settlement cycles (12 cycles/day) against a bank statement, with 3-pass validation."
    version = "1.0.0"
    input_slots = [
        InputSlot(
            key="ntsl", label="NTSL Files", accept=[".xls", ".xlsx", ".xlsm", ".xlsb"], multiple=True, required=True,
            filename_pattern=r"UPI_NTSLERT\d{6}_\w+\.(xls|xlsx|xlsm|xlsb)", help_text="Up to 500 files.",
        ),
        InputSlot(key="bank", label="Bank Statement", accept=[".xlsx", ".xls", ".xlsm", ".xlsb", ".csv"], required=True),
    ]
    options_schema = UPIOptions

    def parse(self, ctx: RunContext) -> ParsedData:
        ntsl_files = ctx.files_for("ntsl")
        bank_file = ctx.files_for("bank")[0]

        all_rows: list[dict[str, Any]] = []
        parse_errors: list[dict[str, str]] = []
        for f in ntsl_files:
            try:
                rows = read_rows(f.content, f.filename)
                all_rows.extend(parse_ntsl(rows, f.filename))
            except Exception as exc:  # noqa: BLE001
                parse_errors.append({"file": f.filename, "reason": str(exc)})

        bank_rows = read_rows(bank_file.content, bank_file.filename)
        bank_data = load_bank(bank_rows)

        ctx.info(f"Parsed {len(all_rows)} working rows from {len(ntsl_files) - len(parse_errors)} files, {len(bank_data)} bank entries")
        return ParsedData(
            data={"ntsl_files": ntsl_files, "all_rows": all_rows, "parse_errors": parse_errors, "bank_data": bank_data, "file_count": len(ntsl_files)}
        )

    def validate(self, parsed: ParsedData, ctx: RunContext) -> list[ValidationFinding]:
        d = parsed.data
        findings = _run_pass1(d["ntsl_files"], bank_present=True)
        if d["parse_errors"]:
            for pe in d["parse_errors"][:8]:
                findings.append(ValidationFinding(pass_name="Pass 1 — Input Integrity", check=f"Parse error — {pe['file']}", level="warn", message=pe["reason"]))

        ntsl_net = _aggregate_cycle_nets(d["all_rows"])[0]
        findings.extend(_run_pass2(d["all_rows"], ntsl_net, d["file_count"]))
        return findings

    def reconcile(self, parsed: ParsedData, ctx: RunContext, findings: list[ValidationFinding]) -> ReconOutput:
        if any(f.level == "err" for f in findings):
            return ReconOutput(stats=[StatCard(label="Status", value="Failed pre-checks", tone="bad")])

        all_rows = parsed.data["all_rows"]
        bank_data = parsed.data["bank_data"]
        ntsl_net, ntsl_detail = _aggregate_cycle_nets(all_rows)

        all_ntsl_dates = sorted({k.split("_")[0] for k in ntsl_net})
        min_ntsl_date = all_ntsl_dates[0] if all_ntsl_dates else ""

        match_res: list[dict[str, Any]] = []
        for b in bank_data:
            b_net = b["bankCredit"] - b["bankDebit"]
            if not b["ntslDate"]:
                match_res.append({**b, "ntslNet": None, "ntslDebit": None, "ntslCredit": None, "diff": None, "matchStatus": "Parse Error", "npciRef": None})
                continue
            dk = date_key(b["ntslDate"])
            k = f"{dk}_{b['ntslCycle']}"
            nd = ntsl_detail.get(k)
            nv = ntsl_net.get(k)
            is_prior = dk < min_ntsl_date
            if nv is None:
                status = "Prior Month Carryover" if is_prior else "Not Found in NTSL"
                diff_val = None
            else:
                diff_val = round2(b_net - nv)
                if is_prior:
                    status = "Matched (Prior Month)" if abs(diff_val) < MATCH_TOL else "Discrepancy (Prior Month)"
                else:
                    status = "Matched" if abs(diff_val) < MATCH_TOL else "Discrepancy"
            if nd:
                nd["status"] = status
            match_res.append(
                {
                    **b, "ntslNet": nd["ntslNet"] if nd else None, "ntslDebit": nd["ntslDebit"] if nd else None,
                    "ntslCredit": nd["ntslCredit"] if nd else None, "diff": diff_val, "matchStatus": status,
                    "npciRef": dk.replace("-", "") + "_" + b["ntslCycle"],
                }
            )

        pass3_findings, act_disc = _run_pass3(bank_data, ntsl_net, match_res)

        apr_mr = [r for r in match_res if "Prior" not in (r["matchStatus"] or "")]
        pr_mr = [r for r in match_res if "Prior" in (r["matchStatus"] or "")]
        matched = sum(1 for r in apr_mr if r["matchStatus"] == "Matched")
        ntsl_total_net = round2(sum(r["ntslNet"] or 0 for r in apr_mr))

        stats = [
            StatCard(label="NTSL Files", value=str(parsed.data["file_count"])),
            StatCard(label="Working Rows", value=str(len(all_rows))),
            StatCard(label="Settlement Cycles", value=str(len(ntsl_detail))),
            StatCard(label="Bank Entries", value=str(len(match_res))),
            StatCard(label="Matched", value=str(matched), tone="good"),
            StatCard(label="Discrepancies", value=str(len(act_disc)), tone="bad" if act_disc else "good"),
            StatCard(label="Prior Month Carryover", value=str(len(pr_mr)), tone="warn" if pr_mr else "neutral"),
            StatCard(label="NTSL Net Settlement", value=f"₹{ntsl_total_net:,.2f}"),
        ]

        settlement_columns = ["NTSL Date", "Cycle", "Settlement Basis", "Debit", "Credit", "Net Settlement", "Status"]
        settlement_rows = []
        for k in sorted(ntsl_detail.keys(), key=lambda x: (x.split("_")[0], CYCLE_ORDER.index(x.split("_")[1]) if x.split("_")[1] in CYCLE_ORDER else 99)):
            nd = ntsl_detail[k]
            cyc = k.split("_")[1]
            settlement_rows.append(
                {
                    "NTSL Date": nd["date"].isoformat(), "Cycle": cyc, "Settlement Basis": "T" if cyc in T_CYCLES else "T+1",
                    "Debit": nd["ntslDebit"], "Credit": nd["ntslCredit"], "Net Settlement": nd["ntslNet"], "Status": nd["status"],
                }
            )

        bank_columns = ["NTSL Date", "Cycle", "Settlement Basis", "NTSL Net", "Bank Debit", "Bank Credit", "Bank Net", "Difference", "Match Status", "Narrative"]
        bank_rows = []
        for r in match_res:
            cyc = r.get("ntslCycle") or ""
            bank_rows.append(
                {
                    "NTSL Date": r["ntslDate"].isoformat() if r["ntslDate"] else None, "Cycle": cyc,
                    "Settlement Basis": "T" if cyc in T_CYCLES else ("T+1" if cyc in T1_CYCLES else ""),
                    "NTSL Net": r["ntslNet"], "Bank Debit": r["bankDebit"], "Bank Credit": r["bankCredit"],
                    "Bank Net": round2(r["bankCredit"] - r["bankDebit"]), "Difference": r["diff"], "Match Status": r["matchStatus"],
                    "Narrative": r["narrative"],
                }
            )

        day_map: dict[str, dict[str, float]] = {}
        for r in all_rows:
            dk = date_key(r["date"])
            acc = day_map.setdefault(dk, {"t_deb": 0.0, "t_cred": 0.0, "t1_deb": 0.0, "t1_cred": 0.0})
            if r["cycle"] in T_CYCLES:
                acc["t_deb"] += r["debit"]
                acc["t_cred"] += r["credit"]
            else:
                acc["t1_deb"] += r["debit"]
                acc["t1_cred"] += r["credit"]
        daily_columns = ["Date", "T Debit (3C-9C)", "T Credit (3C-9C)", "T Net", "T+1 Debit", "T+1 Credit", "T+1 Net", "Daily Net"]
        daily_rows = []
        for dk in sorted(day_map.keys()):
            v = day_map[dk]
            t_net = round2(v["t_cred"] - v["t_deb"])
            t1_net = round2(v["t1_cred"] - v["t1_deb"])
            daily_rows.append(
                {
                    "Date": dk, "T Debit (3C-9C)": round2(v["t_deb"]), "T Credit (3C-9C)": round2(v["t_cred"]), "T Net": t_net,
                    "T+1 Debit": round2(v["t1_deb"]), "T+1 Credit": round2(v["t1_cred"]), "T+1 Net": t1_net,
                    "Daily Net": round2(t_net + t1_net),
                }
            )

        cat_map: dict[str, dict[str, float]] = {}
        for r in all_rows:
            cat = r["remarks"] or "Unknown"
            acc = cat_map.setdefault(cat, {"debit": 0.0, "credit": 0.0})
            acc["debit"] += r["debit"]
            acc["credit"] += r["credit"]
        pivot_columns = ["Category", "Sum of Debit", "Sum of Credit", "Net (Credit-Debit)"]
        ordered_cats = [c for c in REMARKS_ORDER if c in cat_map] + sorted(c for c in cat_map if c not in REMARKS_ORDER)
        pivot_rows = [
            {"Category": c, "Sum of Debit": round2(cat_map[c]["debit"]), "Sum of Credit": round2(cat_map[c]["credit"]),
             "Net (Credit-Debit)": round2(cat_map[c]["credit"] - cat_map[c]["debit"])}
            for c in ordered_cats
        ]

        exception_rows = [r for r in bank_rows if r["Match Status"] not in ("Matched",) and "Prior" not in r["Match Status"]]

        sheets = [
            ReportSheet(name="Working", kind="detail", columns=["Date", "Cycle", "Description", "No of Txns", "Debit", "Credit", "Remarks"],
                        rows=[{"Date": r["date"].isoformat(), "Cycle": r["cycle"], "Description": r["desc"], "No of Txns": r["noTxns"], "Debit": r["debit"], "Credit": r["credit"], "Remarks": r["remarks"]} for r in all_rows]),
            ReportSheet(name="Settlement Recon", kind="detail", columns=settlement_columns, rows=settlement_rows),
            ReportSheet(name="Bank Recon", kind="detail", columns=bank_columns, rows=bank_rows),
            ReportSheet(name="Daily Summary", kind="summary", columns=daily_columns, rows=daily_rows),
            ReportSheet(name="Pivot", kind="summary", columns=pivot_columns, rows=pivot_rows),
            ReportSheet(name="Exceptions", kind="exception", columns=bank_columns, rows=exception_rows),
        ]

        return ReconOutput(
            stats=stats, sheets=sheets, findings=pass3_findings,
            raw={"bank_columns": bank_columns, "bank_rows": bank_rows, "settlement_columns": settlement_columns,
                 "settlement_rows": settlement_rows, "daily_columns": daily_columns, "daily_rows": daily_rows,
                 "pivot_columns": pivot_columns, "pivot_rows": pivot_rows,
                 "working_rows": all_rows},
        )

    def build_report(self, output: ReconOutput, ctx: RunContext) -> bytes:
        wb = Workbook()
        raw = output.raw

        def sheet_from_rows(ws_title: str, columns: list[str], rows: list[dict[str, Any]], header_fill: str = "FF1F4E79", is_active: bool = False) -> None:
            ws = wb.active if is_active else wb.create_sheet(ws_title)
            if is_active:
                ws.title = ws_title
            for i, name in enumerate(columns, start=1):
                set_cell(ws, 1, i, name, fill=solid_fill(header_fill))
                bold_font(ws.cell(row=1, column=i), color="FFFFFFFF")
            freeze(ws, "A2")
            for r_i, row in enumerate(rows, start=2):
                write_row(ws, r_i, [row.get(c) for c in columns])
            set_column_widths(ws, {i: 16 for i in range(1, len(columns) + 1)})

        working_rows = [{"Date": r["date"].isoformat(), "Cycle": r["cycle"], "Description": r["desc"], "No of Txns": r["noTxns"], "Debit": r["debit"], "Credit": r["credit"], "Remarks": r["remarks"]} for r in raw["working_rows"]]
        sheet_from_rows("Working", ["Date", "Cycle", "Description", "No of Txns", "Debit", "Credit", "Remarks"], working_rows, is_active=True)
        sheet_from_rows("Settlement Recon", raw["settlement_columns"], raw["settlement_rows"])
        sheet_from_rows("Bank Recon", raw["bank_columns"], raw["bank_rows"])
        sheet_from_rows("Daily Summary", raw["daily_columns"], raw["daily_rows"])
        sheet_from_rows("Pivot", raw["pivot_columns"], raw["pivot_rows"])

        insights = wb.create_sheet("Insights")
        set_cell(insights, 1, 1, f"{COMPANY} — UPI Settlement Insights", fill=solid_fill("FF1F4E79"))
        bold_font(insights.cell(row=1, column=1), color="FFFFFFFF")
        r_i = 3
        for stat in output.stats:
            set_cell(insights, r_i, 1, stat.label)
            set_cell(insights, r_i, 2, stat.value)
            r_i += 1
        set_column_widths(insights, {1: 28, 2: 22})

        import io

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _aggregate_cycle_nets(all_rows: list[dict[str, Any]]) -> tuple[dict[str, float], dict[str, dict[str, Any]]]:
    cagg: dict[str, dict[str, Any]] = {}
    for r in all_rows:
        k = date_key(r["date"]) + "_" + r["cycle"]
        acc = cagg.setdefault(k, {"date": r["date"], "d": 0.0, "c": 0.0, "td": 0.0, "tc": 0.0})
        acc["d"] += r["debit"]
        acc["c"] += r["credit"]
        if r["remarks"] == "Transaction Amount":
            acc["td"] += r["debit"]
            acc["tc"] += r["credit"]

    ntsl_net: dict[str, float] = {}
    ntsl_detail: dict[str, dict[str, Any]] = {}
    for k, v in cagg.items():
        ntsl_net[k] = round4(v["c"] - v["d"])
        ntsl_detail[k] = {
            "date": v["date"], "ntslDebit": round2(v["d"]), "ntslCredit": round2(v["c"]),
            "ntslNet": round2(v["c"] - v["d"]), "txnDebit": round2(v["td"]), "txnCredit": round2(v["tc"]), "status": "Pending",
        }
    return ntsl_net, ntsl_detail


register(UPIModule())
