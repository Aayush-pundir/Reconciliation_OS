"""Payment Gateway Recon (Ops vs PayU vs Cashfree) — faithful port of
`ReconModule.bas` (`RunReconciliation` / `DetermineAction`).

The rule table below preserves the VBA's exact sequential if/return order
(first matching rule wins) - including its acknowledged-redundant rules
and its fallback - because re-ordering or "simplifying" it would silently
change which action wins for status triples that satisfy multiple rules.
"""
from __future__ import annotations

from typing import Any

from app.recon.base import (
    InputSlot,
    OptionsBase,
    ParsedData,
    ReconModule,
    ReconOutput,
    ReportSheet,
    RunContext,
    StatCard,
    ValidationFinding,
)
from app.recon.registry import register
from app.recon.utils.columns import find_col
from app.recon.utils.excel_read import read_rows, rows_to_dicts
from app.recon.utils.excel_write import bold_font, freeze, set_cell, solid_fill, write_row, set_column_widths
from openpyxl import Workbook

ACTION_COLORS = {
    "Successfully Completed": ("FFC6EFCE", "FF1E7B1E"),
    "Refund": ("FFFFEB9C", "FF9C5700"),
    "Failed": ("FFFFC7CE", "FF9C0006"),
    "Cancelled": ("FFE0C2CD", "FF6B2E42"),
    "To Check": ("FFFFD966", "FF6B4E00"),
    "Manual Check Required": ("FFF4CCCC", "FF990000"),
}
HEADER_FILL = "FF1F4E79"


class PGReconOptions(OptionsBase):
    pass


def normalize_key(value: Any) -> str:
    s = str(value or "").strip().upper()
    return s if s.startswith("TRX") else ""


def determine_action(ops: str, pyu: str, cf: str) -> tuple[str, str]:
    """Returns (action, remark). Mirrors `DetermineAction` in ReconModule.bas
    rule-for-rule, in the exact same order."""
    pyu_nf = pyu in ("not found", "")
    cf_nf = cf in ("NOT FOUND", "")

    if ops == "FAILED" and pyu == "failed" and cf_nf:
        return "Failed", ""
    if ops in ("FAILED", "CANCELED") and pyu == "captured" and cf_nf:
        return "Refund", "PayU shows captured but Ops is failed/canceled — initiate refund"
    if ops == "FAILED" and pyu == "dropped" and cf_nf:
        return "Failed", ""
    if ops in ("FAILED", "CANCELED") and pyu_nf and cf == "SUCCESS":
        return "Refund", "Cashfree SUCCESS but Ops failed/canceled — verify and process refund"
    if ops in ("FAILED", "CANCELED") and pyu_nf and cf_nf:
        return "Failed", ""
    if ops in ("FAILED", "CANCELED") and pyu_nf and cf == "FAILED":
        return "Failed", ""
    if ops == "CANCELED" and (pyu_nf or pyu == "failed") and (cf_nf or cf == "FAILED"):
        return "Cancelled", ""
    if ops == "FAILED" and pyu == "usercancelled" and cf_nf:
        return "Failed", ""
    if ops == "COMPLETED" and pyu == "captured" and cf_nf:
        return "Successfully Completed", ""
    if ops == "COMPLETED" and pyu_nf and cf == "SUCCESS":
        return "Successfully Completed", ""
    if ops in ("MANUAL", "INITIATED", "EXCEPTION") and (
        pyu == "failed" or pyu == "captured" or cf == "FAILED" or cf == "SUCCESS"
    ):
        return "To Check", "Manual/Initiated/Exception state — verify with payment gateway"
    if ops == "COMPLETED" and pyu == "captured" and cf == "SUCCESS":
        return "Successfully Completed", ""
    if ops == "COMPLETED" and (pyu == "captured" or pyu_nf) and (cf == "SUCCESS" or cf_nf):
        return "Successfully Completed", ""
    if ops in ("MANUAL", "INITIATED", "EXCEPTION"):
        return "To Check", "Ops status requires manual verification"
    if ops in ("FAILED", "CANCELED") and (pyu == "captured" or cf == "SUCCESS"):
        return "Refund", "Payment gateway shows success but Ops failed — refund likely required"
    if (
        ops in ("FAILED", "CANCELED")
        and (pyu_nf or pyu in ("failed", "dropped", "usercancelled"))
        and (cf_nf or cf == "FAILED")
    ):
        return "Failed", ""
    return "Manual Check Required", "Unexpected status combination — manual review required"


class PGReconModule(ReconModule):
    key = "pg_recon"
    display_name = "Payment Gateway Recon (Ops / PayU / Cashfree)"
    description = "Cross-references Ops transaction status against PayU and Cashfree gateway records."
    version = "1.0.0"
    input_slots = [
        InputSlot(key="ops", label="Ops (Add Money)", accept=[".xlsx", ".xls", ".csv"], required=True),
        InputSlot(key="payu", label="PayU", accept=[".xlsx", ".xls", ".csv"], required=True),
        InputSlot(key="cashfree", label="Cashfree", accept=[".xlsx", ".xls", ".csv"], required=True),
    ]
    options_schema = PGReconOptions

    def parse(self, ctx: RunContext) -> ParsedData:
        ops_file = ctx.files_for("ops")[0]
        payu_file = ctx.files_for("payu")[0]
        cf_file = ctx.files_for("cashfree")[0]

        ops = rows_to_dicts(read_rows(ops_file.content, ops_file.filename))
        payu = rows_to_dicts(read_rows(payu_file.content, payu_file.filename))
        cashfree = rows_to_dicts(read_rows(cf_file.content, cf_file.filename))
        ctx.info(f"Ops: {len(ops)} rows | PayU: {len(payu)} rows | Cashfree: {len(cashfree)} rows")
        return ParsedData(data={"ops": ops, "payu": payu, "cashfree": cashfree})

    def validate(self, parsed: ParsedData, ctx: RunContext) -> list[ValidationFinding]:
        findings: list[ValidationFinding] = []
        checks = [
            ("ops", ["Transaction ID", "Status"]),
            ("payu", ["txnid", "status"]),
            ("cashfree", ["Order Id", "Transaction Status"]),
        ]
        for slot, required in checks:
            rows = parsed.data[slot]
            if not rows:
                findings.append(
                    ValidationFinding(pass_name="Input Integrity", check=slot, level="err", message="No rows found")
                )
                continue
            headers = list(rows[0].keys())
            missing = [r for r in required if find_col(headers, r) is None]
            if missing:
                findings.append(
                    ValidationFinding(
                        pass_name="Input Integrity",
                        check=f"{slot} headers",
                        level="err",
                        message=f"Missing required column(s): {', '.join(missing)}",
                    )
                )
            else:
                findings.append(
                    ValidationFinding(pass_name="Input Integrity", check=f"{slot} headers", level="ok", message="OK")
                )
        return findings

    def reconcile(self, parsed: ParsedData, ctx: RunContext, findings: list[ValidationFinding]) -> ReconOutput:
        if any(f.level == "err" for f in findings):
            return ReconOutput(stats=[StatCard(label="Status", value="Failed pre-checks", tone="bad")])

        ops_rows = parsed.data["ops"]
        payu_rows = parsed.data["payu"]
        cf_rows = parsed.data["cashfree"]

        ops_headers = list(ops_rows[0].keys())
        payu_headers = list(payu_rows[0].keys())
        cf_headers = list(cf_rows[0].keys())

        c_ops_txn = find_col(ops_headers, "Transaction ID")
        c_ops_stat = find_col(ops_headers, "Status")
        c_payu_txn = find_col(payu_headers, "txnid")
        c_payu_stat = find_col(payu_headers, "status")
        c_cf_txn = find_col(cf_headers, "Order Id")
        c_cf_stat = find_col(cf_headers, "Transaction Status")

        payu_status: dict[str, str] = {}
        for row in payu_rows:
            key = normalize_key(row.get(c_payu_txn))
            if key:
                payu_status[key] = str(row.get(c_payu_stat) or "").strip()

        cf_status: dict[str, str] = {}
        for row in cf_rows:
            key = normalize_key(row.get(c_cf_txn))
            if key:
                cf_status[key] = str(row.get(c_cf_stat) or "").strip()

        detail_rows: list[dict[str, Any]] = []
        processed_ops_keys: set[str] = set()
        action_counts: dict[str, int] = {}
        coverage_counts = {"Ops + PayU + CF": 0, "Ops + PayU only": 0, "Ops + CF only": 0, "Ops only": 0}

        for row in ops_rows:
            key = normalize_key(row.get(c_ops_txn))
            if not key:
                continue
            processed_ops_keys.add(key)
            ops_status = str(row.get(c_ops_stat) or "").strip().upper()
            in_payu = key in payu_status
            in_cf = key in cf_status
            pyu = payu_status[key].lower().strip() if in_payu else "not found"
            cf = cf_status[key].upper().strip() if in_cf else "NOT FOUND"

            action, remark = determine_action(ops_status, pyu, cf)
            coverage = (
                "Ops + PayU + CF"
                if in_payu and in_cf
                else "Ops + PayU only"
                if in_payu
                else "Ops + CF only"
                if in_cf
                else "Ops only"
            )
            coverage_counts[coverage] += 1
            action_counts[action] = action_counts.get(action, 0) + 1
            detail_rows.append(
                {
                    "Transaction ID": key,
                    "Ops Status": ops_status,
                    "PayU Status": "Not Found" if pyu == "not found" else pyu,
                    "Cashfree Status": "Not Found" if cf == "NOT FOUND" else cf,
                    "Action Point": action,
                    "Remarks": remark,
                    "Source Coverage": coverage,
                }
            )

        payu_only = cf_only = 0
        for key, status in payu_status.items():
            if key in processed_ops_keys:
                continue
            payu_only += 1
            action_counts["Manual Check Required"] = action_counts.get("Manual Check Required", 0) + 1
            cf_disp = cf_status.get(key, "Not Found").upper() if key in cf_status else "Not Found"
            detail_rows.append(
                {
                    "Transaction ID": key,
                    "Ops Status": "NOT IN OPS",
                    "PayU Status": status.lower(),
                    "Cashfree Status": cf_disp,
                    "Action Point": "Manual Check Required",
                    "Remarks": "TXN present in PayU but missing from Ops sheet",
                    "Source Coverage": "PayU only",
                }
            )

        for key, status in cf_status.items():
            if key in processed_ops_keys or key in payu_status:
                continue
            cf_only += 1
            action_counts["Manual Check Required"] = action_counts.get("Manual Check Required", 0) + 1
            detail_rows.append(
                {
                    "Transaction ID": key,
                    "Ops Status": "NOT IN OPS",
                    "PayU Status": "Not Found",
                    "Cashfree Status": status.upper(),
                    "Action Point": "Manual Check Required",
                    "Remarks": "TXN present in Cashfree only — not found in Ops or PayU",
                    "Source Coverage": "CF only",
                }
            )

        total = len(detail_rows)
        columns = [
            "Transaction ID",
            "Ops Status",
            "PayU Status",
            "Cashfree Status",
            "Action Point",
            "Remarks",
            "Source Coverage",
        ]
        exception_rows = [r for r in detail_rows if r["Action Point"] != "Successfully Completed"]

        stats = [StatCard(label="Total Transactions", value=str(total), tone="neutral")]
        for action in [
            "Successfully Completed",
            "Refund",
            "Failed",
            "Cancelled",
            "To Check",
            "Manual Check Required",
        ]:
            count = action_counts.get(action, 0)
            tone = "good" if action == "Successfully Completed" else ("bad" if count else "neutral")
            pct = f"{(count / total * 100):.1f}%" if total else "0.0%"
            stats.append(StatCard(label=action, value=str(count), tone=tone, sub=pct))
        for label, count in coverage_counts.items():
            stats.append(StatCard(label=label, value=str(count)))
        stats.append(StatCard(label="PayU only (missing from Ops)", value=str(payu_only)))
        stats.append(StatCard(label="Cashfree only (missing from Ops)", value=str(cf_only)))

        sheets = [
            ReportSheet(name="Reconciliation", kind="detail", columns=columns, rows=detail_rows),
            ReportSheet(name="Exceptions", kind="exception", columns=columns, rows=exception_rows),
        ]
        return ReconOutput(stats=stats, sheets=sheets, findings=[], raw={"detail_rows": detail_rows, "columns": columns})

    def build_report(self, output: ReconOutput, ctx: RunContext) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Reconciliation"
        columns = output.raw["columns"]
        set_column_widths(ws, {1: 20, 2: 14, 3: 14, 4: 16, 5: 22, 6: 45, 7: 18})
        for i, name in enumerate(columns, start=1):
            set_cell(ws, 1, i, name, fill=solid_fill(HEADER_FILL))
            bold_font(ws.cell(row=1, column=i), color="FFFFFFFF")
        freeze(ws, "A2")

        row_idx = 2
        for row in output.raw["detail_rows"]:
            bg, fg = ACTION_COLORS.get(row["Action Point"], ("FFFFFFFF", "FF000000"))
            write_row(ws, row_idx, [row[c] for c in columns], fill=solid_fill(bg))
            bold_font(ws.cell(row=row_idx, column=5), color=fg)
            row_idx += 1

        summary = wb.create_sheet("Summary")
        set_cell(summary, 1, 1, "PAYMENT RECONCILIATION SUMMARY", fill=solid_fill(HEADER_FILL))
        bold_font(summary.cell(row=1, column=1), color="FFFFFFFF")
        summary.merge_cells("A1:C1")
        r = 3
        set_cell(summary, r, 1, "Action Point")
        set_cell(summary, r, 2, "Count")
        set_cell(summary, r, 3, "% of Total")
        for stat in output.stats:
            r += 1
            set_cell(summary, r, 1, stat.label)
            set_cell(summary, r, 2, stat.value)
            set_cell(summary, r, 3, stat.sub or "")
        set_column_widths(summary, {1: 34, 2: 12, 3: 12})

        import io

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


register(PGReconModule())
