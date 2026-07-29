"""6F/6R Data Preparation — port of the Streamlit `app.py` (latest version:
`EXCLUDE_TOKEN = "reversal"`, all-sheets scan, Normal/Gifted/NCMC source
labels). See PRD §7.7.

Difference from the original tool: the original took server-local folder
paths (`st.text_input` of a Windows path) because it ran on an analyst's
own machine. Recon OS runs server-side behind a browser client, which has
no access to a local folder path - so the input model here is direct
multi-file upload (the file *type* classification and every processing
rule below is unchanged).
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from app.recon.base import (
    InputSlot,
    OptionsBase,
    ParsedData,
    ReconModule,
    ReconOutput,
    ReportSheet,
    RunContext,
    StatCard,
    ValidationFinding,
)
from app.recon.registry import register
from app.recon.utils.excel_write import bold_font, freeze, set_cell, set_column_widths, solid_fill, write_row
from openpyxl import Workbook

import pandas as pd

LARGE_FILE_MB_THRESHOLD = 10
EXCLUDE_TOKEN = "reversal"

try:
    import python_calamine  # noqa: F401

    ENGINE_KW: dict[str, str] = {"engine": "calamine"}
except ImportError:
    ENGINE_KW = {}

REQUIRED_COLS = {
    "NORMAL": {"Status", "TransactionCategory", "TransactionDate", "TransactionType", "Amount"},
    "GIFT": {"Status", "TransactionCategory", "TransactionDate", "TransactionType", "Amount"},
    "NCMC": {"SETTLEMENTDATE", "SETTLEMENTAMT", "TRXN_TYPE"},
}
SOURCE_LABEL = {"NORMAL": "Normal", "GIFT": "Gifted", "NCMC": "NCMC"}


class DataPrep6FOptions(OptionsBase):
    pass


def detect_file_type(filename: str, size_bytes: int) -> str:
    name = filename.upper()
    size_mb = size_bytes / (1024 * 1024)
    if name.startswith("GIFT"):
        return "GIFT"
    if name.startswith("NCMC"):
        return "NCMC"
    if size_mb > LARGE_FILE_MB_THRESHOLD:
        return "NORMAL"
    return "UNKNOWN"


def read_all_relevant_sheets(content: bytes, filename: str, ftype: str) -> list[pd.DataFrame]:
    required = REQUIRED_COLS.get(ftype, set())
    if filename.lower().endswith(".csv"):
        df = pd.read_csv(io.BytesIO(content))
        return [df] if required.issubset(set(df.columns)) else []

    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, **ENGINE_KW)
    return [df for df in sheets.values() if required.issubset(set(df.columns))]


def to_ddmmyyyy(series: "pd.Series", source: str) -> "pd.Series":
    if source == "NCMC":
        return pd.to_datetime(series.astype(str), format="%Y%m%d", errors="coerce").dt.strftime("%d-%m-%Y")
    parsed = pd.to_datetime(series, errors="coerce", dayfirst=True)
    return parsed.dt.strftime("%d-%m-%Y")


def process_normal(df: "pd.DataFrame") -> "pd.DataFrame":
    df = df.copy()
    df["Status"] = df["Status"].astype(str).str.strip().str.lower()
    df = df[df["Status"] == "completed"]
    cat = df["TransactionCategory"].fillna("").astype(str).str.strip().str.lower()
    df = df[~cat.str.contains(EXCLUDE_TOKEN, na=False)]
    df["__date__"] = to_ddmmyyyy(df["TransactionDate"], "NORMAL")
    df["__type__"] = df["TransactionType"].astype(str).str.strip().str.upper()
    df["__amount__"] = pd.to_numeric(df["Amount"], errors="coerce").fillna(0)
    return df[["__date__", "__type__", "__amount__"]]


def process_gift(df: "pd.DataFrame") -> "pd.DataFrame":
    df = df.copy()
    df["Status"] = df["Status"].astype(str).str.strip().str.lower()
    df = df[df["Status"] == "completed"]
    cat = df["TransactionCategory"].fillna("").astype(str).str.strip().str.lower()
    df = df[~cat.str.contains(EXCLUDE_TOKEN, na=False)]
    df["__date__"] = to_ddmmyyyy(df["TransactionDate"], "GIFT")
    df["__type__"] = df["TransactionType"].astype(str).str.strip().str.upper()
    df["__amount__"] = pd.to_numeric(df["Amount"], errors="coerce").fillna(0)
    return df[["__date__", "__type__", "__amount__"]]


def process_ncmc(df: "pd.DataFrame") -> "pd.DataFrame":
    df = df.copy()
    df["__date__"] = to_ddmmyyyy(df["SETTLEMENTDATE"], "NCMC")
    df["__type__"] = df["TRXN_TYPE"].astype(str).str.strip().str.upper()
    df["__amount__"] = pd.to_numeric(df["SETTLEMENTAMT"], errors="coerce").fillna(0)
    return df[["__date__", "__type__", "__amount__"]]


PROCESSORS = {"NORMAL": process_normal, "GIFT": process_gift, "NCMC": process_ncmc}
EXPECTED_COLS = ["Credit-Normal", "Credit-NCMC", "Credit-Gifted", "Debit-Normal", "Debit-NCMC", "Debit-Gifted"]


def aggregate(tagged_frames: dict[str, list["pd.DataFrame"]]) -> "pd.DataFrame":
    pieces = []
    for source, frames in tagged_frames.items():
        if not frames:
            continue
        df = pd.concat(frames, ignore_index=True)
        if df.empty:
            continue
        g = df.groupby(["__date__", "__type__"])["__amount__"].sum().reset_index()
        g["__source__"] = source
        pieces.append(g)

    if not pieces:
        return pd.DataFrame(columns=["Date", *EXPECTED_COLS])

    combined = pd.concat(pieces, ignore_index=True)

    def type_norm(t: str) -> str:
        t = t.upper()
        if "CREDIT" in t or t == "CR":
            return "Credit"
        if "DEBIT" in t or t == "DR":
            return "Debit"
        return t.title()

    combined["__type_norm__"] = combined["__type__"].apply(type_norm)
    combined["__source_label__"] = combined["__source__"].map(SOURCE_LABEL).fillna(combined["__source__"])
    combined["__col__"] = combined["__type_norm__"] + "-" + combined["__source_label__"]

    pivot = combined.pivot_table(
        index="__date__", columns="__col__", values="__amount__", aggfunc="sum", fill_value=0
    ).reset_index().rename(columns={"__date__": "Date"})

    for col in EXPECTED_COLS:
        if col not in pivot.columns:
            pivot[col] = 0
    pivot = pivot[["Date", *EXPECTED_COLS]]

    pivot["__sort__"] = pd.to_datetime(pivot["Date"], format="%d-%m-%Y", errors="coerce")
    pivot = pivot.sort_values("__sort__").drop(columns="__sort__").reset_index(drop=True)
    return pivot


class DataPrep6FModule(ReconModule):
    key = "data_prep_6f"
    display_name = "6F/6R Data Preparation"
    description = "Classifies Normal/GIFT/NCMC transaction extracts and builds a daily Credit/Debit aggregate report."
    version = "1.0.0"
    input_slots = [
        InputSlot(
            key="files",
            label="Transaction Files",
            accept=[".csv", ".xlsx", ".xls"],
            multiple=True,
            required=True,
            help_text="GIFT*/NCMC* files classified by filename prefix; other files >10MB classified as Normal.",
        ),
    ]
    options_schema = DataPrep6FOptions

    def parse(self, ctx: RunContext) -> ParsedData:
        files = ctx.files_for("files")
        tagged: dict[str, list[pd.DataFrame]] = {"NORMAL": [], "GIFT": [], "NCMC": []}
        skipped: list[dict[str, str]] = []

        for f in files:
            ftype = detect_file_type(f.filename, f.size_bytes)
            if ftype == "UNKNOWN":
                skipped.append({"file": f.filename, "reason": "could not classify (not GIFT/NCMC prefix, not >10MB)"})
                continue
            try:
                frames = read_all_relevant_sheets(f.content, f.filename, ftype)
                if not frames:
                    skipped.append({"file": f.filename, "reason": "no relevant sheet found (required columns missing)"})
                    continue
                for raw in frames:
                    tagged[ftype].append(PROCESSORS[ftype](raw))
            except Exception as exc:  # noqa: BLE001
                skipped.append({"file": f.filename, "reason": str(exc)})

        ctx.info(f"Processed {len(files)} file(s), {len(skipped)} skipped")
        return ParsedData(data={"tagged": tagged, "skipped": skipped, "file_count": len(files)})

    def validate(self, parsed: ParsedData, ctx: RunContext) -> list[ValidationFinding]:
        findings: list[ValidationFinding] = []
        tagged = parsed.data["tagged"]
        total_frames = sum(len(v) for v in tagged.values())
        if total_frames == 0:
            findings.append(
                ValidationFinding(
                    pass_name="Input Integrity",
                    check="Classified sheets",
                    level="err",
                    message="No files could be classified into NORMAL/GIFT/NCMC with required columns present",
                )
            )
        else:
            findings.append(
                ValidationFinding(
                    pass_name="Input Integrity",
                    check="Classified sheets",
                    level="ok",
                    message=f"{total_frames} relevant sheet(s) found across NORMAL/GIFT/NCMC",
                )
            )
        if parsed.data["skipped"]:
            findings.append(
                ValidationFinding(
                    pass_name="Input Integrity",
                    check="Skipped files",
                    level="warn",
                    message=f"{len(parsed.data['skipped'])} file(s) skipped — see Exceptions tab",
                )
            )
        return findings

    def reconcile(self, parsed: ParsedData, ctx: RunContext, findings: list[ValidationFinding]) -> ReconOutput:
        tagged = parsed.data["tagged"]
        skipped = parsed.data["skipped"]
        report = aggregate(tagged)

        columns = ["Date", *EXPECTED_COLS]
        rows = [
            {c: (None if pd.isna(v) else (float(v) if c != "Date" else v)) for c, v in zip(columns, r)}
            for r in report.itertuples(index=False, name=None)
        ]

        total_credit = sum(r.get(c) or 0 for r in rows for c in ("Credit-Normal", "Credit-NCMC", "Credit-Gifted"))
        total_debit = sum(r.get(c) or 0 for r in rows for c in ("Debit-Normal", "Debit-NCMC", "Debit-Gifted"))

        stats = [
            StatCard(label="Files Processed", value=str(parsed.data["file_count"])),
            StatCard(label="Files Skipped", value=str(len(skipped)), tone="warn" if skipped else "good"),
            StatCard(label="Report Rows (Dates)", value=str(len(rows))),
            StatCard(label="Total Credit", value=f"₹{total_credit:,.2f}", tone="good"),
            StatCard(label="Total Debit", value=f"₹{total_debit:,.2f}"),
        ]

        sheets = [
            ReportSheet(name="6F Report", kind="detail", columns=columns, rows=rows),
            ReportSheet(
                name="Exceptions",
                kind="exception",
                columns=["file", "reason"],
                rows=skipped,
            ),
        ]
        return ReconOutput(stats=stats, sheets=sheets, findings=[], raw={"columns": columns, "rows": rows})

    def build_report(self, output: ReconOutput, ctx: RunContext) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "6F_Report"
        columns = output.raw["columns"]
        set_column_widths(ws, {i: 16 for i in range(1, len(columns) + 1)} | {1: 14})
        for i, name in enumerate(columns, start=1):
            set_cell(ws, 1, i, name, fill=solid_fill("FF1F3864"))
            bold_font(ws.cell(row=1, column=i), color="FFFFFFFF")
        freeze(ws, "A2")

        row_idx = 2
        for row in output.raw["rows"]:
            values = [row.get(c) for c in columns]
            write_row(ws, row_idx, values)
            for col_idx in range(2, len(columns) + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = "#,##0.00"
            row_idx += 1

        import io as _io

        buf = _io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


register(DataPrep6FModule())
