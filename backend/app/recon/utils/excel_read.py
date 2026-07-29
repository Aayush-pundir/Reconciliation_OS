"""Unified workbook reader: xls / xlsx / xlsm / xlsb / csv / HTML-as-XLS.

Every source tool had to content-sniff files because filenames/extensions
lie (NPCI's ".xls" exports are actually HTML). We do the same sniffing
server-side (by magic bytes), where it can't be bypassed by a spoofed
Content-Type header - see ARCHITECTURE.md §10 Security Notes.
"""
from __future__ import annotations

import io
from typing import Any

import openpyxl
import pandas as pd

from app.recon.utils.html_xls import looks_like_html, parse_html_table

ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0"


def sniff_format(data: bytes) -> str:
    if data[:4] == ZIP_MAGIC:
        return "ooxml"
    if data[:4] == OLE_MAGIC:
        return "ole"
    if looks_like_html(data):
        return "html"
    return "unknown"


def _rows_from_openpyxl(data: bytes, sheet_name: str | None = None) -> list[list[Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _rows_from_pandas(data: bytes, filename: str) -> list[list[Any]]:
    df = pd.read_excel(io.BytesIO(data), header=None, dtype=object, sheet_name=0)
    return df.where(pd.notnull(df), None).values.tolist()


def read_rows(data: bytes, filename: str = "") -> list[list[Any]]:
    """Read the first sheet of any supported format into row-major cells
    (list[list[Any]], no header row stripped) - the shape every module's
    port of the original `header:1` parsing expects.
    """
    if filename.lower().endswith(".csv"):
        df = pd.read_csv(io.BytesIO(data), header=None, dtype=object)
        return df.where(pd.notnull(df), None).values.tolist()

    fmt = sniff_format(data)
    if fmt == "html":
        return parse_html_table(data)
    try:
        return _rows_from_openpyxl(data)
    except Exception:
        return _rows_from_pandas(data, filename)


def read_all_sheets(data: bytes, filename: str = "") -> dict[str, list[list[Any]]]:
    """Read every sheet - required by the 6F module, which scans all sheets
    of a workbook for the presence of its required-column set."""
    if filename.lower().endswith(".csv") or sniff_format(data) == "html":
        return {"Sheet1": read_rows(data, filename)}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        return {name: [list(row) for row in wb[name].iter_rows(values_only=True)] for name in wb.sheetnames}
    except Exception:
        sheets = pd.read_excel(io.BytesIO(data), header=None, dtype=object, sheet_name=None)
        return {name: df.where(pd.notnull(df), None).values.tolist() for name, df in sheets.items()}


def rows_to_dicts(rows: list[list[Any]], header_row_index: int = 0) -> list[dict[str, Any]]:
    """Convenience: turn row-major cells into header-keyed dicts, used by
    modules that prefer named-column access after locating the header row."""
    if header_row_index >= len(rows):
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[header_row_index]]
    out: list[dict[str, Any]] = []
    for row in rows[header_row_index + 1 :]:
        out.append({headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))})
    return out
